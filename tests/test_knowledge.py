"""Actual persistence/parsers and controlled provider boundaries; no paid calls."""

import io

from phase2_support import (
    kb,
    private_settings,  # noqa: F401
    upload,
)


def test_knowledge_lifecycle_and_formats(platform):
    c, app, root, ts, hs, users, password = platform
    a, b = hs
    kid = kb(c, a)
    from docx import Document as Docx
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    word = io.BytesIO()
    d = Docx()
    d.add_heading("Warranty", 1)
    d.add_paragraph("Warranty lasts thirty days.")
    d.add_table(rows=1, cols=1).cell(0, 0).text = "Support by email."
    d.save(word)
    pdf = io.BytesIO()
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 100 Td (Warranty lasts thirty days.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(pdf)
    docs = []
    for filename, content in [
        ("policy.txt", b"Warranty lasts thirty days."),
        ("guide.md", b"# Warranty\n\nWarranty lasts thirty days."),
        ("manual.docx", word.getvalue()),
        ("manual.pdf", pdf.getvalue()),
    ]:
        r = upload(c, a, kid, filename, content)
        assert r.status_code == 201, r.text
        assert r.json()["status"] == "ready", r.text
        assert r.json()["chunk_count"] > 0 and len(r.json()["sha256"]) == 64
        docs.append(r.json())
    repeated = upload(c, a, kid, "renamed.txt", b"Warranty lasts thirty days.")
    assert repeated.status_code == 200 and repeated.json()["id"] == docs[0]["id"]
    found = c.post(f"/api/v1/knowledge-bases/{kid}/search", headers=a, json={"query": "Warranty"})
    assert found.status_code == 200 and len(found.json()) == 4, found.text
    assert any(h["page"] == 1 for h in found.json())
    assert any(h["section"] == "Warranty" for h in found.json())
    assert (
        c.post(
            f"/api/v1/knowledge-bases/{kid}/search", headers=a, json={"query": "unfindableword"}
        ).json()
        == []
    )
    did = docs[0]["id"]
    for method, path in [
        ("get", f"/documents/{did}"),
        ("post", f"/documents/{did}/reprocess"),
        ("delete", f"/documents/{did}"),
        ("get", f"/knowledge-bases/{kid}"),
        ("delete", f"/knowledge-bases/{kid}"),
    ]:
        assert getattr(c, method)("/api/v1" + path, headers=b).status_code == (
            204 if method == "delete" else 404
        )
    assert (
        c.post(
            f"/api/v1/knowledge-bases/{kid}/search", headers=b, json={"query": "Warranty"}
        ).status_code
        == 404
    )
    assert c.post(f"/api/v1/documents/{did}/reprocess", headers=a).json()["status"] == "ready"
    assert c.delete(f"/api/v1/documents/{did}", headers=a).status_code == 204
    assert c.delete(f"/api/v1/documents/{did}", headers=a).status_code == 204
    hits = c.post(
        f"/api/v1/knowledge-bases/{kid}/search", headers=a, json={"query": "Warranty"}
    ).json()
    assert not any(h["document_id"] == did for h in hits)
    assert not app.state.storage.exists(ts[0]["id"], did)
    assert upload(c, a, kid, "bad.exe", b"not text").status_code == 422
    assert upload(c, a, kid, "../bad.txt", b"bad").status_code == 422
    assert upload(c, a, kid, "bad.txt", b"\x00\xff").status_code == 422
    assert upload(c, a, kid, "bad.pdf", b"plain text").status_code == 422
    damaged = upload(c, a, kid, "damaged.pdf", b"%PDF-1.7\ninvalid")
    assert damaged.status_code == 201 and damaged.json()["status"] == "failed"
    assert "Traceback" not in damaged.text
    assert c.delete(f"/api/v1/knowledge-bases/{kid}", headers=a).status_code == 204
    assert c.get(f"/api/v1/knowledge-bases/{kid}/documents", headers=a).status_code == 404
    assert all(not app.state.storage.exists(ts[0]["id"], doc["id"]) for doc in docs)


def test_concurrent_delete_and_locked_row_refresh(platform, monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    from types import SimpleNamespace
    from uuid import UUID

    from support_agent.db import set_tenant
    from support_agent.models import Document
    from support_agent.services import knowledge
    from support_agent.services.resources import get_owned

    c, app, root, ts, hs, users, password = platform
    a = hs[0]
    kid = kb(c, a)
    did = upload(c, a, kid, "race.txt", b"Concurrent deletion evidence.").json()["id"]
    # A previously read ORM object must be refreshed after waiting for a row lock.
    context = SimpleNamespace(tenant_id=UUID(ts[0]["id"]), public_session_id=None)
    with app.state.sessions() as first:
        set_tenant(first, context.tenant_id)
        old = first.get(Document, UUID(did))
        assert old.status == "ready"
        with app.state.sessions.begin() as second:
            set_tenant(second, context.tenant_id)
            current = second.get(Document, UUID(did))
            current.status = "failed"
        assert get_owned(first, Document, context, UUID(did), lock=True).status == "failed"
        first.rollback()
    original = knowledge.resume_context
    barrier = Barrier(2)

    def synchronize(db, ctx, request, **kwargs):
        result = original(db, ctx, request, **kwargs)
        if request.method == "DELETE" and not getattr(request.state, "barrier_passed", False):
            request.state.barrier_passed = True
            barrier.wait(timeout=5)
        return result

    monkeypatch.setattr(knowledge, "resume_context", synchronize)
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(
            pool.map(
                lambda _: c.delete(f"/api/v1/documents/{did}", headers=a).status_code, range(2)
            )
        )
    assert results == [204, 204]
    assert not app.state.storage.exists(ts[0]["id"], did)


def test_storage_cleanup_retry_and_deleted_evidence(platform, monkeypatch):
    c, app, root, ts, hs, users, password = platform
    a = hs[0]
    kid = kb(c, a)
    did = upload(c, a, kid, "cleanup.txt", b"Cleanup evidence survives no deletion.").json()["id"]
    original = app.state.storage.delete

    def unavailable(*args):
        raise OSError("synthetic-secret-do-not-log")

    monkeypatch.setattr(app.state.storage, "delete", unavailable)
    response = c.delete(f"/api/v1/documents/{did}", headers=a)
    assert response.status_code == 503
    events = c.get("/api/v1/audit-events", headers=a).json()
    assert any(
        event["resource_id"] == did and event["reason"] == "storage_cleanup_pending"
        for event in events
    )
    assert c.get(f"/api/v1/documents/{did}", headers=a).json()["status"] == "deleting"
    assert (
        c.post(f"/api/v1/knowledge-bases/{kid}/search", headers=a, json={"query": "Cleanup"}).json()
        == []
    )
    monkeypatch.setattr(app.state.storage, "delete", original)
    assert c.delete(f"/api/v1/documents/{did}", headers=a).status_code == 204
    assert not app.state.storage.exists(ts[0]["id"], did)


def test_upload_and_parser_resource_limits(platform, monkeypatch):
    import zipfile
    from types import SimpleNamespace

    from support_agent import ingestion

    c, app, root, ts, hs, users, password = platform
    a = hs[0]
    kid = kb(c, a)
    assert upload(c, a, kid, "broken.docx", b"PK\x03\x04invalid ZIP").status_code == 422
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "placeholder")
        z.writestr("word/document.xml", "x" * (9 * 1024 * 1024))
    assert upload(c, a, kid, "expanded.docx", bomb.getvalue()).status_code == 422
    assert upload(c, a, kid, "large.txt", b"x" * (6 * 1024 * 1024)).status_code == 413
    monkeypatch.setattr(
        ingestion.psutil,
        "Process",
        lambda pid: SimpleNamespace(memory_info=lambda: SimpleNamespace(rss=513 * 1024 * 1024)),
    )
    chunks, error = ingestion.process_document(b"Resource limit probe.", ".txt")
    assert chunks is None and error == "parse_resource_limit"


def test_failed_reprocess_hides_old_chunks(platform, monkeypatch):
    from support_agent.services import knowledge

    c, app, root, ts, hs, users, password = platform
    a = hs[0]
    kid = kb(c, a)
    did = upload(c, a, kid, "replacement.txt", b"Replacement policy.").json()["id"]
    assert c.post(
        f"/api/v1/knowledge-bases/{kid}/search", headers=a, json={"query": "Replacement"}
    ).json()
    monkeypatch.setattr(knowledge, "process_document", lambda *args: (None, "parse_failed"))
    response = c.post(f"/api/v1/documents/{did}/reprocess", headers=a)
    assert response.json()["status"] == "failed" and response.json()["chunk_count"] == 0
    assert (
        c.post(
            f"/api/v1/knowledge-bases/{kid}/search", headers=a, json={"query": "Replacement"}
        ).json()
        == []
    )
